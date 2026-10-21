"""Write learned outputs into a copy of the input, read them back, and compare. Formula caches are kept."""

from datetime import date, datetime
from decimal import Decimal
import os
from pathlib import Path
import posixpath
import tempfile
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook

from .sources import bind_tables, canon, equivalent, find_tables, sheet_rows


def plain(value):
    """JSON value that keeps the type visible, so 1 and "1" stay distinguishable in feedback."""
    if isinstance(value, dict):
        return {k: plain(v) for k, v in value.items()}
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value if value is None or isinstance(value, (bool, int, float, str)) else str(value)


def output_spec(spec):
    """Only the output tables and their columns, for reading a target or a written workbook."""
    return {"outputs": spec["outputs"], "tables": {key: {**spec["tables"][key], "columns": {
        name: spec["tables"][key]["columns"][name] for name in out["columns"]}} for key, out in spec["outputs"].items()}}


def read_outputs(path, spec):
    """Output tables of a workbook: appended rows as a list, filled columns by record key."""
    narrowed = output_spec(spec)
    bound, problems = bind_tables(narrowed, find_tables(sheet_rows(path)))
    if problems:
        raise ValueError("Output tables not found in " + Path(path).name + ": " + "; ".join(p["question"] for p in problems))
    result = {}
    for key, out in spec["outputs"].items():
        table = bound[key]
        if out["op"] == "append_rows":
            result[key] = [{c: r["values"][c] for c in out["columns"]} for r in table["all"]]
        else:
            result[key] = {canon(r["key"]) or r["row"]: {c: r["values"][c] for c in out["columns"]} for r in table["records"]}
    return result


def compare_outputs(actual, expected):
    differences = []
    for key, want in expected.items():
        got = actual.get(key, [] if isinstance(want, list) else {})
        if isinstance(want, list):
            if len(got) != len(want):
                differences.append({"table": key, "check": "row_count", "actual": len(got), "expected": len(want)})
            pairs = list(enumerate(zip(got, want), 1))
        else:
            if set(got) != set(want):
                differences.append({"table": key, "check": "records", "actual": sorted(map(str, got)), "expected": sorted(map(str, want))})
            pairs = [(record, (got[record], want[record])) for record in want if record in got]
        for row, (g, w) in pairs:
            for column in w:
                if not equivalent(g.get(column), w[column]):
                    differences.append({"table": key, "row": row, "column": column,
                                        "actual": plain(g.get(column)), "expected": plain(w[column])})
    return differences


def put(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # never let a text value starting with "=" become a formula
    elif isinstance(value, date) and cell.number_format == "General":
        cell.number_format = "yyyy-mm-dd"


def write_outputs(source, destination, bound, outputs):
    """A new workbook: appended tables written under their header, filled columns at their records.

    Appending needs empty rows under the header; existing rows stop the run instead of being deleted.
    """
    destination = Path(destination)
    if destination.exists() or destination.resolve() == Path(source).resolve():
        raise ValueError("Output must be a new path; original workbooks are never overwritten.")
    book = load_workbook(source, data_only=False)
    try:
        for key, out in outputs.items():
            table = bound[key]
            if table.get("file") not in (None, str(source)):
                raise ValueError(f"{key} was found in {Path(table['file']).name}; outputs are only written into the input workbook.")
            sheet = book[table["sheet"]]
            columns = table["columns"]
            if out["op"] == "append_rows":
                first = table["row"] + 1
                taken = [r for r in range(first, first + len(out["rows"]))
                         if any(sheet.cell(r, c + 1).value is not None for c in columns.values())]
                if taken:
                    raise ValueError(f"{key} must be empty under its header, but rows {taken[:5]} already hold data; nothing was written.")
                for offset, row in enumerate(out["rows"], first):
                    for name, value in row.items():
                        put(sheet.cell(offset, columns[name] + 1), value)
                continue
            for ctx, row in zip(out["contexts"], out["rows"]):
                for name, value in row.items():
                    put(sheet.cell(ctx["record"]["row"], columns[name] + 1), value)
            # A summary row of =SUM formulas right after the records gets totals for the filled numbers.
            total = next((p for p in table["post"] if any(isinstance(f, str) and f.upper().startswith("=SUM(")
                                                          for f in p["formulas"])), None)
            if total:
                for name in out["rows"][0] if out["rows"] else []:
                    numbers = [row[name] for row in out["rows"]]
                    cell = sheet.cell(total["row"], columns[name] + 1)
                    if cell.value is None and all(isinstance(v, (int, float, Decimal)) and not isinstance(v, bool) for v in numbers):
                        cell.value = sum(Decimal(str(v)) for v in numbers)
        with tempfile.TemporaryDirectory(dir=destination.parent) as directory:
            raw, ready = Path(directory) / "raw.xlsx", Path(directory) / "ready.xlsx"
            book.save(raw)
            restore_caches(source, raw, ready)
            os.link(ready, destination)
    finally:
        book.close()
    return destination


def sheet_parts(archive):
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {r.attrib["Id"]: r.attrib["Target"] for r in rels}
    book = ET.fromstring(archive.read("xl/workbook.xml"))
    result = {}
    for sheet in book.find(f"{{{main}}}sheets"):
        target = targets[sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]]
        result[sheet.attrib["name"]] = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
    return result


def restore_caches(original, generated, destination):
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    tag = lambda name: f"{{{main}}}{name}"
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(generated) as edited:
        old_parts, new_parts = sheet_parts(source), sheet_parts(edited)
        replacements = {}
        for name in old_parts.keys() & new_parts.keys():
            old = ET.fromstring(source.read(old_parts[name]))
            new = ET.fromstring(edited.read(new_parts[name]))
            cached = {cell.attrib["r"]: cell for cell in old.iter(tag("c")) if cell.find(tag("f")) is not None}
            for cell in new.iter(tag("c")):
                prior = cached.get(cell.attrib["r"])
                formula = cell.find(tag("f"))
                if prior is None or formula is None or formula.text != prior.find(tag("f")).text:
                    continue
                value = prior.find(tag("v"))
                if value is None or value.text is None:
                    continue
                current = cell.find(tag("v"))
                if current is None:
                    current = ET.SubElement(cell, tag("v"))
                current.text = value.text
                if "t" in prior.attrib:
                    cell.attrib["t"] = prior.attrib["t"]
                else:
                    cell.attrib.pop("t", None)
            replacements[new_parts[name]] = ET.tostring(new)
        with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as result:
            for item in edited.infolist():
                result.writestr(item, replacements.get(item.filename, edited.read(item.filename)))
