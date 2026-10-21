"""Phase 3: update an existing workflow when input headers change. Model proposes, code verifies."""

from datetime import date, datetime
import json
from pathlib import Path
import re

from openpyxl import load_workbook

from .events import Events
from .planner import run_workflow
from .sources import meaning, normalise
from .workbook import HEADERS, alias_sets, learned_headers

GROUPS = {"roster": "roster", "tracker": "roster", "constants": "constants", "fund": "constants",
          "accounts": "accounts", "staging": "staging"}
MEANINGS = {
    "roster.id": ("investor ID", {"text"}), "roster.investor": ("investor name", {"text"}),
    "roster.called": ("amount called from the investor this round", {"number", "currency"}),
    "tracker.received": ("amount actually received from the investor", {"number", "currency"}),
    "tracker.outstanding": ("called minus received", {"number", "currency"}),
    "tracker.receipt_date": ("date the payment arrived per the bank statement", {"date"}),
    "accounts.transaction_type": ("transaction type name in the chart of accounts", {"text"}),
    "accounts.account": ("account code in the chart of accounts", {"text", "number"}),
    "constants.entity": ("fund legal entity", {"text"}), "constants.bank_account": ("fund bank account", {"text"}),
    "constants.batch_type": ("journal batch type", {"text"}), "fund.entity": ("fund legal entity", {"text"}),
    "fund.entity_domain": ("legal entity domain", {"text"}), "fund.deal": ("deal", {"text"}),
    "fund.position": ("position", {"text"}), "fund.batch_type": ("journal batch type", {"text"}),
    "fund.bank_account": ("fund bank account", {"text"}),
    "staging.date": ("bank transaction date", None), "staging.description": ("bank reference text", None),
    "staging.withdrawal": ("withdrawal amount", None), "staging.deposit": ("deposit amount", None),
    "staging.investor": ("matched investor name", None)}
PROMPT_VERSION = "adapt-headers-v1"
SYSTEM = """A saved spreadsheet workflow no longer finds some input columns because the
input file changed (headers renamed, columns moved). For each missing field, choose the
candidate header with the same business meaning, using the field meaning, the previous
header and the value types under each candidate. Return null for a field when no
candidate clearly has that meaning, or when two candidates could both be it; then
explain in question what a person must confirm. Never choose by position alone.
Return JSON: {"mappings": {field: header or null}, "question": text or null}."""


def recognised(name, text):
    if not isinstance(text, str) or not normalise(text):
        return False
    if name == "roster.called" and meaning(text) == "called":
        return True
    return any(normalise(text) in names for names in alias_sets(name))


def kind(value):
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    if isinstance(value, str):
        if value.startswith("="):
            return "formula"
        return "currency" if re.fullmatch(r"[A-Z]{3}", value.strip()) else "text"
    return None


def header_rows(path):
    """Rows with at least two text cells, plus the value types beneath each column."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = []
        for sheet in book:
            # headers are assumed within the first 200 rows of a sheet; widen if real files bury them deeper.
            grid = [list(row) for row in sheet.iter_rows(max_row=200, values_only=True)]
            for index, row in enumerate(grid):
                cells = {column: value for column, value in enumerate(row) if isinstance(value, str) and value.strip()}
                if len(cells) < 2 or set(HEADERS) <= set(cells.values()):
                    continue  # the DIU upload template is output, not a source table
                below = {column: sorted({kind(r[column]) for r in grid[index + 1:index + 16]
                                         if column < len(r) and kind(r[column])}) for column in cells}
                rows.append({"sheet": sheet.title, "row": index + 1, "cells": cells, "below": below})
        return rows
    finally:
        book.close()


def missing_fields(headers, rows):
    """Known fields whose header cannot be found in the best-matching header row of their table."""
    result = {}
    for group in set(GROUPS.values()):
        names = [name for name in headers if GROUPS[name.split(".")[0]] == group]
        if not names:
            continue
        best = max(rows, key=lambda row: sum(any(recognised(n, text) for n in names) for text in row["cells"].values()),
                   default=None)
        if best is None:
            continue
        lost = [n for n in names if not any(recognised(n, text) for text in best["cells"].values())]
        claimed = {c for c, text in best["cells"].items() if any(recognised(n, text) for n in headers)}
        candidates = [{"header": text, "values": best["below"][c]} for c, text in best["cells"].items() if c not in claimed]
        for name in lost:
            result[name] = {"table": f"{best['sheet']}!row {best['row']}", "candidates": candidates}
    return result


def proposal_schema(missing):
    return {"type": "object", "additionalProperties": False, "required": ["mappings", "question"],
            "properties": {
                "mappings": {"type": "object", "additionalProperties": False, "required": list(missing),
                             "properties": {name: {"anyOf": [{"type": "string", "enum": [c["header"] for c in info["candidates"]]},
                                                             {"type": "null"}]} for name, info in missing.items()}},
                "question": {"type": ["string", "null"]}}}


def empty(candidate):
    return set(candidate["values"]) <= {"currency"}


def verify(mapping, missing, answered=()):
    """Code-side checks on a proposed mapping: a real candidate, used once, compatible values.

    Fields the workflow writes (tracker/staging) must map to an empty column, and when more than one
    empty column fits, the values give no evidence, so only a person's answer can choose.
    """
    problems = []
    chosen = [header for header in mapping.values() if header is not None]
    if len(chosen) != len(set(chosen)):
        problems.append("The same header was chosen for two different fields.")
    for name, header in mapping.items():
        if header is None:
            continue
        candidate = next((c for c in missing[name]["candidates"] if c["header"] == header), None)
        if candidate is None:
            problems.append(f"{name}: '{header}' is not a candidate header.")
            continue
        if name.split(".")[0] in {"tracker", "staging"}:
            if not empty(candidate):
                problems.append(f"{name}: '{header}' already holds data; the workflow would overwrite it.")
                continue
            fits = [c["header"] for c in missing[name]["candidates"] if empty(c)]
            if len(fits) > 1 and name not in answered:
                problems.append(f"{name}: {len(fits)} empty columns could hold it ({', '.join(fits)}); "
                                "the values give no evidence, so this needs a person to choose.")
                continue
        allowed = MEANINGS[name][1]
        unexpected = set(candidate["values"]) - (allowed or set(candidate["values"])) - {"formula"}
        if unexpected:
            problems.append(f"{name}: '{header}' holds {sorted(unexpected)} values; expected {sorted(allowed)}.")
    return problems


def adapt_workflow(path, workbook, bank, context, client, directory, answers=None):
    """answers: {field: header} chosen by a person for earlier column_mapping questions."""
    directory = Path(directory)
    saved = json.loads(Path(path).read_text())
    version = saved.get("workflow_version", 1)
    first = run_workflow(path, workbook, bank, directory / "run-current", saved)
    if first["status"] == "success":
        return {**first, "status": "success", "adapted": False, "workflow": str(path),
                "note": "The existing workflow already fits this input; no update needed."}
    headers = saved.get("headers", {})
    with learned_headers(headers):
        missing = missing_fields(headers, header_rows(workbook))
    if not missing:
        return {**first, "adapted": False,
                "note": "No renamed or moved columns detected; the run stopped for the reasons listed."}
    gone = [name for name, info in missing.items() if not info["candidates"]]
    if gone:
        questions = [{"code": "column_missing", "field": name, "question":
                      f"The {MEANINGS[name][0]} column (previously '{headers[name]}') is not in {missing[name]['table']} "
                      "and no unrecognised column is left there. Add it back or say where it moved."} for name in gone]
        (Path(directory) / "questions.json").write_text(json.dumps(questions, indent=2))
        return {"status": "needs_input", "adapted": False, "questions": questions, "run_dir": str(directory)}
    request = {"context": Path(context).read_text() if context else "", "workflow_intent": saved["plan"]["intent"],
               "missing_fields": [{"field": name, "meaning": MEANINGS[name][0], "previous_header": headers[name],
                                   "expected_values": sorted(MEANINGS[name][1] or []) or "any", "table": info["table"]}
                                  for name, info in missing.items()],
               "candidates": {name: info["candidates"] for name, info in missing.items()}}
    (directory / "adapt-request.json").write_text(json.dumps(request, indent=2, default=str))
    answers = {name: header for name, header in (answers or {}).items() if name in missing}
    reply = {"mappings": {}, "question": None}
    ask = {name: info for name, info in missing.items() if name not in answers}
    if ask:
        client.schema = proposal_schema(ask)
        request["missing_fields"] = [f for f in request["missing_fields"] if f["field"] in ask]
        reply = json.loads(client.complete([{"role": "system", "content": SYSTEM},
                                            {"role": "user", "content": json.dumps(request, default=str)}],
                                           Events(directory / "model"), 1))
    mapping = {**{k: v for k, v in reply.get("mappings", {}).items() if k in ask}, **answers}
    unresolved = [name for name in missing if mapping.get(name) is None]
    problems = verify(mapping, missing, answers)
    if unresolved or problems:
        questions = [{"code": "column_mapping", "field": name, "question":
                      f"Which column now holds the {MEANINGS[name][0]} (previously '{headers[name]}')?",
                      "candidates": [c["header"] for c in missing[name]["candidates"]]} for name in unresolved]
        questions += [{"code": "mapping_rejected", "question": problem} for problem in problems]
        # A rejected pick becomes a question a person can answer from the same candidates.
        questions += [{"code": "column_mapping", "field": name, "question": f"Which column now holds the {MEANINGS[name][0]} "
                       f"(previously '{headers[name]}')?", "candidates": [c["header"] for c in missing[name]["candidates"]]}
                      for name in missing if name not in unresolved and any(p.startswith(name + ":") for p in problems)]
        if reply.get("question"):
            questions.append({"code": "model_question", "question": reply["question"]})
        (directory / "questions.json").write_text(json.dumps(questions, indent=2))
        return {"status": "needs_input", "adapted": False, "mapping": mapping, "questions": questions, "run_dir": str(directory)}
    updated = {**saved, "workflow_version": version + 1, "headers": {**headers, **mapping},
               "adaptations": saved.get("adaptations", []) + [{
                   "from_version": version, "input": Path(workbook).name,
                   "mapped": {name: {"from": headers[name], "to": header,
                                     "decided_by": "person" if name in answers else "model, verified by code"}
                              for name, header in mapping.items()}}]}
    check = run_workflow(path, workbook, bank, directory / "run-adapted", updated)
    if check["status"] != "success":
        return {**check, "adapted": False, "mapping": mapping,
                "note": "The proposed mapping was applied but the workflow checks still fail; nothing was saved."}
    destination = directory / "workflow.json"
    destination.write_text(json.dumps(updated, indent=2))
    return {**check, "status": "success", "adapted": True, "mapping": mapping,
            "workflow": str(destination), "workflow_version": version + 1}
