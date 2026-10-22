"""Command line: inspect, build (phase 1), run (phase 2), adapt (phase 3) and the local UI.

When a person is at the terminal, every question is asked there and the command carries on with the answer."""

import argparse
import json
import os
from pathlib import Path
import sys

from openpyxl import load_workbook

from .adapt import PROMPT_VERSION as ADAPT_PROMPT_VERSION, adapt_workflow
from .events import Events, run_name
from .model import OpenRouter, settings
from .builder import MAX_ATTEMPTS, PROMPT_VERSION as BUILD_PROMPT_VERSION, build_workflow
from .planner import run_workflow


def inspect_workbook(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = []
        for sheet in workbook:
            preview = []
            populated_rows = 0
            for row in sheet.iter_rows():
                cells = []
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cell.value
                    if cell.data_type == "f":
                        value = getattr(value, "text", value)
                    cells.append({"cell": cell.coordinate, "value": value})
                if cells:
                    populated_rows += 1
                    if len(preview) < 8:
                        preview.append(cells)
            sheets.append({
                "name": sheet.title,
                "reported_rows": sheet.max_row,
                "reported_columns": sheet.max_column,
                "populated_rows": populated_rows,
                "preview": preview,
            })
        return {"file": str(path), "sheets": sheets}
    finally:
        workbook.close()


def prompt(questions, reader=sys.stdin, writer=sys.stderr):
    """Ask a person each question on the terminal: pick a numbered option or type any answer; Enter skips."""
    answers = {}
    for item in questions:
        writer.write(f"\n? {item['question']}\n")
        options = item.get("options") or []
        for number, option in enumerate(options, 1):
            writer.write(f"  {number}. {option}\n")
        writer.write(("  Choose a number or type an answer" if options else "  Type an answer") + " (Enter to skip)\n> ")
        writer.flush()
        reply = reader.readline().strip()
        if not reply:
            continue
        if reply.isdigit() and 1 <= int(reply) <= len(options):
            reply = options[int(reply) - 1]
        answers[item["id"]] = {"question": item["question"], "answer": reply}
    return answers


def run_until_settled(args, answers, ask):
    """Phase 2 has no model, so a question stops it; when a person is there, ask and run again."""
    asked, directory, round_ = set(), args.run_dir, 1
    result = run_workflow(args.workflow, args.input, args.sources, directory, answers=answers)
    while result["status"] == "needs_input" and ask:
        fresh = [q for q in result["questions"] if q["id"] not in asked]
        if not fresh:
            break
        asked.update(q["id"] for q in fresh)
        answers.update(ask(fresh))
        round_ += 1
        directory = args.run_dir.with_name(f"{args.run_dir.name}-{round_}")
        result = run_workflow(args.workflow, args.input, args.sources, directory, answers=answers)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Build, run and adapt a spreadsheet workflow.")
    commands = parser.add_subparsers(dest="command", required=True)
    inspect = commands.add_parser("inspect", help="Read XLSX structure and a small preview.")
    inspect.add_argument("workbook", type=Path)
    full = commands.add_parser("build", help="Phase 1: learn, execute and verify a workflow against a target.")
    full.add_argument("--input", type=Path, required=True)
    full.add_argument("--target", type=Path, required=True)
    full.add_argument("--sources", type=Path, nargs="*", default=[], help="Other source files: workbooks or PDF tables (e.g. a bank statement).")
    full.add_argument("--context", type=Path, required=True)
    full.add_argument("--max-attempts", type=int, choices=range(1, MAX_ATTEMPTS + 1), default=MAX_ATTEMPTS)
    run = commands.add_parser("run", help="Phase 2: run a saved workflow on an input; no target or model.")
    run.add_argument("--workflow", type=Path, required=True)
    run.add_argument("--input", type=Path, required=True)
    run.add_argument("--sources", type=Path, nargs="*", default=[], help="Other source files: workbooks or PDF tables (e.g. a bank statement).")
    adapt = commands.add_parser("adapt", help="Phase 3: update a saved workflow for a changed input; saves v+1.")
    adapt.add_argument("--workflow", type=Path, required=True)
    adapt.add_argument("--input", type=Path, required=True)
    adapt.add_argument("--sources", type=Path, nargs="*", default=[], help="Other source files: workbooks or PDF tables (e.g. a bank statement).")
    adapt.add_argument("--context", type=Path)
    for command in (full, run, adapt):
        command.add_argument("--answers", type=Path, help='JSON {"question id": "answer"}, e.g. from an earlier questions.json.')
        command.add_argument("--no-input", action="store_true", help="Never prompt; stop and write questions.json instead.")
        command.add_argument("--run-dir", type=Path)
    commands.add_parser("ui", help="Open the local web UI.")
    args = parser.parse_args()
    if args.command == "ui":
        script = Path(__file__).with_name("ui.py")
        os.execvp(sys.executable, [sys.executable, "-m", "streamlit", "run", str(script)])
    if args.command != "inspect" and args.run_dir is None:
        args.run_dir = Path("runs") / run_name(args.command)
    try:
        if args.command == "inspect":
            result = inspect_workbook(args.workbook)
        else:
            answers = json.loads(args.answers.read_text()) if args.answers else {}
            ask = None if args.no_input or not sys.stdin.isatty() else prompt
            if args.command == "build":
                key, model = settings()
                result = build_workflow(args.input, args.target, args.sources, args.context, answers,
                                        OpenRouter(key, model, BUILD_PROMPT_VERSION, 8000),
                                        Events(args.run_dir, (key,)), args.max_attempts, ask)
            elif args.command == "adapt":
                key, model = settings()
                result = adapt_workflow(args.workflow, args.input, args.sources, args.context,
                                        OpenRouter(key, model, ADAPT_PROMPT_VERSION, 1500), args.run_dir, answers, ask)
            else:
                result = run_until_settled(args, answers, ask)
    except Exception as error:
        parser.exit(1, f"Command failed: {error}\n")
    print(json.dumps({k: v for k, v in result.items() if k != "bindings"}, indent=2, default=str))
    if result.get("status") in {"needs_input", "budget_exhausted", "stopped"}:
        parser.exit(2)


if __name__ == "__main__":
    main()
