import json
from pathlib import Path
from tempfile import TemporaryDirectory
import time
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from ylookup.planner import run_workflow
from ylookup.sources import LEAVE
from fixtures import built, edit, pages, statement


class RunTest(unittest.TestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.folder = Path(directory.name)
        self.workflow, self.source, self.bank, _ = built(self.folder)

    def run_it(self, name, workflow=None, answers=None):
        with patch("ylookup.sources.pdf_pages", return_value=pages()):
            started = time.perf_counter()
            result = run_workflow(workflow or self.workflow, self.source, [self.bank], self.folder / name, answers=answers)
        self.elapsed = time.perf_counter() - started
        return result

    def test_moved_tab_and_columns_are_found_and_reported(self):
        book = load_workbook(self.source)
        book["Fund"].title = "Static"
        for row in (1, 2):
            sheet = book["Static"]
            sheet[f"A{row}"].value, sheet[f"C{row}"].value = sheet[f"C{row}"].value, sheet[f"A{row}"].value
        book.save(self.source)
        result = self.run_it("moved")
        self.assertEqual(result["status"], "success")
        self.assertIn("Fund: tab 'Fund' -> 'Static'", result["binding_changes"])
        self.assertIn("Fund.Legal Entity: column A -> C", result["binding_changes"])

    def test_missing_value_is_a_question_with_a_stable_id_and_an_answer_fills_it(self):
        edit(self.source, ("Roster", "B4", None))
        asked = self.run_it("blank")
        self.assertEqual(asked["status"], "needs_input")
        self.assertEqual([q["id"] for q in asked["questions"]], ["missing_value:Roster.Investor row 4"])
        self.assertFalse((self.folder / "blank/output.xlsx").exists())
        result = self.run_it("answered", answers={"missing_value:Roster.Investor row 4": "Morgan Vale"})
        self.assertEqual(result["status"], "success")
        book = load_workbook(result["output"], data_only=True)
        self.assertEqual(book["Import"]["L5"].value, "Morgan Vale")
        book.close()
        self.assertIn("Morgan Vale", (self.folder / "answered/answers.json").read_text())

    def test_an_unclear_payer_is_asked_about_never_guessed(self):
        for name, roster_change, description in [("nobody", None, "REF UNKNOWN PAYER"),
                                                 ("two", ("Roster", "B4", "Taylor Reeves"), "REF TAYLOR REE")]:
            with self.subTest(name):
                if roster_change:
                    edit(self.source, roster_change)
                with patch("ylookup.sources.pdf_pages", return_value=pages(statement(description=description))):
                    asked = run_workflow(self.workflow, self.source, [self.bank], self.folder / name)
                self.assertEqual(asked["status"], "needs_input")
                link = [q for q in asked["questions"] if q["code"] == "link"]
                self.assertEqual(len(link), 1)
                self.assertIn(LEAVE, link[0]["options"])
                self.assertFalse((self.folder / name / "output.xlsx").exists())
                # Leaving it for follow-up is a valid answer: the learned checks excuse that row's link columns.
                with patch("ylookup.sources.pdf_pages", return_value=pages(statement(description=description))):
                    left = run_workflow(self.workflow, self.source, [self.bank], self.folder / f"{name}-left",
                                        answers={link[0]["id"]: LEAVE})
                self.assertEqual(left["status"], "success", left.get("check_failures"))
                book = load_workbook(left["output"], data_only=True)
                self.assertEqual([c.value for c in book["Staging"][2]], [book["Staging"]["A2"].value, book["Staging"]["B2"].value, 95, None])
                self.assertEqual(book["Import"]["H3"].value, 0)  # nobody was credited with the unlinked payment
                book.close()

    def test_an_output_table_that_already_has_rows_stops_instead_of_being_overwritten(self):
        edit(self.source, ("Import", "A2", "an earlier upload"))
        result = self.run_it("taken")
        self.assertEqual(result["status"], "stopped")
        self.assertIn("must be empty", result["reason"])
        self.assertFalse((self.folder / "taken/output.xlsx").exists())

    def test_learned_checks_stop_a_plan_that_breaks_a_target_convention(self):
        saved = json.loads(self.workflow.read_text())
        saved["plan"]["steps"][3]["columns"]["Debit"] = {"line": ["N", "Y"]}
        broken = self.folder / "broken.json"
        broken.write_text(json.dumps(saved))
        result = self.run_it("broken", broken)
        self.assertEqual(result["status"], "stopped")
        self.assertIn("Import.Debit (line 1) must be 'Y'", result["reason"])

    def test_a_wrong_column_mapping_fails_fast(self):
        edit(self.source, ("Roster", "C1", "Called now"), ("Roster", "G1", "Commitment"), ("Roster", "G3", 1000),
             ("Roster", "G4", 2000), ("Roster", "A30000", "stray formatting far below"))
        asked = self.run_it("renamed")
        self.assertEqual([q["id"] for q in asked["questions"]], ["column_mapping:Roster.Amount Called"])
        self.assertEqual(asked["questions"][0]["options"], ["Called now", "Commitment"])
        rejected = self.run_it("text", answers={"column_mapping:Roster.Amount Called": "Investor"})
        self.assertIn("rejected", rejected["questions"][0]["question"])
        wrong = self.run_it("wrong", answers={"column_mapping:Roster.Amount Called": "Commitment"})
        self.assertEqual(wrong["status"], "success")  # a person's choice is theirs; it still finishes in time
        self.assertLess(self.elapsed, 5)


if __name__ == "__main__":
    unittest.main()
