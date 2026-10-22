"""Rules stated in the customer's description (a rate, a difference) become verified candidates."""
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from ylookup.builder import build_workflow, learn
from ylookup.events import Events
from ylookup.planner import run_workflow
from fixtures import FakeModel, payroll

PLAN = {"intent": "Payslips", "steps": [{"op": "append_rows", "intent": "One payslip per employee", "table": "Payslip",
        "for_each": "Staff", "only_linked": False, "lines": 1, "columns": {
            "Employee": {"col": "Staff.Name"}, "Gross": {"col": "Staff.Gross Pay"},
            "Tax": {"mul": [{"col": "Staff.Gross Pay"}, 0.2]}, "Net": {"sub": [{"col": "Staff.Gross Pay"}, {"out": "Tax"}]}}}]}


class DescriptionRuleTest(unittest.TestCase):
    def test_a_rate_and_a_difference_from_the_description_are_verified_candidates(self):
        with TemporaryDirectory() as folder:
            source, target, notes = payroll(Path(folder))
            learned = learn(source, target, [], {}, notes.read_text())
        columns = learned["profile"]["outputs"][0]["columns"]
        self.assertIn({"mul": [{"col": "Staff.Gross Pay"}, 0.2]}, columns["Tax"]["all rows"])
        self.assertIn({"sub": [{"col": "Staff.Gross Pay"}, {"out": "Tax"}]}, columns["Net"]["all rows"])
        self.assertEqual(learned["questions"], [])  # not mistaken for values from a missing file

    def test_without_the_rule_the_column_is_asked_about_and_a_named_column_is_not(self):
        with TemporaryDirectory() as folder:
            source, target, _ = payroll(Path(folder))
            silent = learn(source, target, [], {}, "I make the payslips.")
            named = learn(source, target, [], {}, "I make the payslips and work out the tax.")
        self.assertEqual([q["columns"] for q in silent["questions"]], [{"Payslip": ["Tax", "Net"]}])  # a cycle, not a rule
        # a column the description names goes to the model with its examples; the unnamed one is still asked about
        self.assertEqual([q["columns"] for q in named["questions"]], [{"Payslip": ["Net"]}])

    def test_the_built_workflow_applies_the_rule_to_new_figures(self):
        with TemporaryDirectory() as folder:
            folder = Path(folder)
            source, target, notes = payroll(folder)
            built = build_workflow(source, target, [], notes, {}, FakeModel([PLAN]), Events(folder / "build"))
            self.assertEqual((built["status"], built["attempts"]), ("success", 1))
            later, _, _ = payroll(folder / "march" if (folder / "march").mkdir() is None else folder, gross=(1200, 4000))
            run = run_workflow(built["workflow"], later, [], folder / "run")
            self.assertEqual(run["status"], "success")
            rows = [[c.value for c in row] for row in load_workbook(run["output"])["Payslip"].iter_rows(min_row=2)]
            self.assertEqual(rows, [["Person 1", 1200, 240, 960], ["Person 2", 4000, 800, 3200]])


if __name__ == "__main__":
    unittest.main()
