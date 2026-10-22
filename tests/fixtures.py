"""A small synthetic case with its own output template, so tests show nothing about the DIU is built in."""

from datetime import datetime
import json

from openpyxl import Workbook, load_workbook


OUTPUT = ["Entity", "Entry", "Line", "Posted", "Type", "Account", "Debit", "Amount", "Comment", "Bank", "Batch", "Partner", "Ref"]
PRIVATE = ["Taylor Reed", "Morgan Vale", "Example fund", "Example account", "2026-02-24", "TAYLOR REE"]


def statement(description="REF TAYLOR REE", currency="USD", kind="deposit", amount="95.00"):
    header = f"{'Date':<15}{'Description':<55}{'Withdrawal':>14}{'Deposit':>14}{'Balance':>14}"
    debit, credit = ("", amount) if kind == "deposit" else (amount, "")
    record = f"{'24 Feb 2026':<15}{description:<55}{debit:>14}{credit:>14}{'900.00':>14}"
    return f"Currency : {currency}\n{header}\n{record}\n   MORE WRAPPED REFERENCE\n\n   Thank you for banking\n"


def pages(*texts, currency="USD"):
    """PDF page texts for the fake bank.pdf (patched in for ylookup.sources.pdf_pages)."""
    return list(texts or [statement(currency=currency)])


def fixtures(folder):
    source, target, bank, context = [folder / name for name in ["input.xlsx", "target.xlsx", "bank.pdf", "context.txt"]]
    book = Workbook()
    roster = book.active
    roster.title = "Roster"
    roster.append(["ID", "Investor", "Amount Called", "Receipt Date", "Amount Received", "Amount outstanding"])
    roster.append([None, None, "USD", None, "USD", "USD"])
    roster.append(["001", "Taylor Reed", 100])
    roster.append(["002", "Morgan Vale", 50])
    roster.append([None, None, "=SUM(C3:C4)"])
    accounts = book.create_sheet("Accounts")
    accounts.append(["Account Code", "Transaction Type"])
    accounts.append(["CASH-CODE", "Cash received"])
    accounts.append(["RECEIVABLE-CODE", "Due from partners"])
    for index in range(60):  # a reference list, like a chart of accounts
        accounts.append([f"FILL-{index}", f"Other movement {index}"])
    fund = book.create_sheet("Fund")
    fund.append(["Legal Entity", "Bank Account", "Batch Type"])
    fund.append(["Example fund", "Example account", "General"])
    stage = book.create_sheet("Staging")
    stage.append(["Date", "Description", "Deposit", "Matched Investor"])
    output = book.create_sheet("Import")
    output.append(OUTPUT)
    book.save(source)
    # The completed example, specified independently of the executor.
    roster.cell(3, 4, datetime(2026, 2, 24))
    roster.cell(3, 5, 95)
    roster.cell(3, 6, 5)
    roster.cell(4, 5, 0)
    roster.cell(4, 6, 50)
    stage.append([datetime(2026, 2, 24), "REF TAYLOR REE MORE WRAPPED REFERENCE", 95, "Taylor Reed"])
    for entry, name, amount, paid in [(1, "Taylor Reed", 95, datetime(2026, 2, 24)), (2, "Morgan Vale", 0, None)]:
        comment = "Paid by " + name
        output.append(["Example fund", entry, 1, paid, "Cash received", "CASH-CODE", "Y", amount, comment, "Example account", "General", None, entry * 2 - 1])
        output.append(["Example fund", entry, 2, paid, "Due from partners", "RECEIVABLE-CODE", "N", amount, comment, None, "General", name, entry * 2])
    book.save(target)
    book.close()
    bank.write_bytes(b"PDF content replaced by a parsed fixture in tests")
    context.write_text("I match payments to partners by name, log what they paid against what was called, "
                       "and build the import with two balancing lines per partner.")
    return source, target, bank, context


def fixture_plan():
    col = lambda table, column: {"col": f"{table}.{column}"}
    var = lambda name: {"var": name}
    line = lambda first, second: {"line": [first, second]}
    return {"intent": "Reconcile receipts and prepare the import", "steps": [
        {"op": "link_rows", "intent": "Identify who paid", "table": "bank", "text": "Description", "to": "Roster",
         "name": "Investor", "policy": "full_or_unique_prefix"},
        {"op": "append_rows", "intent": "Stage the statement", "table": "Staging", "for_each": "bank", "only_linked": False,
         "lines": 1, "columns": {"Date": col("bank", "Date"), "Description": col("bank", "Description"),
                                 "Deposit": col("bank", "Deposit"), "Matched Investor": col("Roster", "Investor")}},
        {"op": "fill_columns", "intent": "Paid versus called", "table": "Roster", "columns": {
            "Receipt Date": {"first": "bank.Date"}, "Amount Received": {"sum": "bank.Deposit"},
            "Amount outstanding": {"sub": [col("Roster", "Amount Called"), {"sum": "bank.Deposit"}]}}},
        {"op": "append_rows", "intent": "Two balancing lines per partner", "table": "Import", "for_each": "Roster",
         "only_linked": False, "lines": 2, "columns": {
             "Entity": col("Fund", "Legal Entity"), "Entry": var("entry"), "Line": var("line"), "Posted": {"first": "bank.Date"},
             "Type": line("Cash received", "Due from partners"),
             "Account": {"lookup": {"table": "Accounts", "where": "Transaction Type", "equals": {"out": "Type"}, "return": "Account Code"}},
             "Debit": line("Y", "N"), "Amount": {"sum": "bank.Deposit"},
             "Comment": {"concat": ["Paid by ", col("Roster", "Investor")]},
             "Bank": line(col("Fund", "Bank Account"), None), "Batch": col("Fund", "Batch Type"),
             "Partner": line(None, col("Roster", "Investor")), "Ref": var("row")}}]}


class FakeModel:
    """Replays canned replies and records every request."""
    model, prompt_version, max_tokens = "test-model", "test", 8000

    def __init__(self, replies):
        self.replies, self.messages, self.schema = iter(replies), [], None

    def complete(self, messages, events, attempt):
        self.messages.append(json.loads(json.dumps(messages)))
        reply = next(self.replies)
        if isinstance(reply, dict) and "steps" in reply:
            reply = {"status": "ready", "workflow": reply, "question": None, "options": None}
        return json.dumps(reply)


def built(folder):
    """The fixture workflow, built through the real builder with a replayed model reply."""
    from pathlib import Path
    from unittest.mock import patch
    from ylookup.builder import build_workflow
    from ylookup.events import Events
    source, target, bank, context = fixtures(folder)
    with patch("ylookup.sources.pdf_pages", return_value=pages()):
        result = build_workflow(source, target, [bank], context, {}, FakeModel([fixture_plan()]), Events(folder / "build"))
    return Path(result["workflow"]), source, bank, context


def edit(path, *changes):
    from openpyxl import load_workbook
    book = load_workbook(path)
    for sheet, cell, value in changes:
        book[sheet][cell] = value
    book.save(path)
    book.close()


PAYROLL = "Each month I make a payslip for every employee: their name and gross pay, tax is 20% of gross, and net is gross minus tax."


def payroll(folder, gross=(1000, 2500, 3200), context=PAYROLL):
    """A payroll job, unrelated to the reconciliation: its rules are only in the description."""
    source, target, notes = folder / "staff.xlsx", folder / "payslips.xlsx", folder / "notes.txt"
    book = Workbook()
    staff = book.active
    staff.title = "Staff"
    staff.append(["Emp ID", "Name", "Gross Pay"])
    for n, amount in enumerate(gross, 1):
        staff.append([f"E{n}", f"Person {n}", amount])
    book.create_sheet("Payslip").append(["Employee", "Gross", "Tax", "Net"])
    book.save(source)
    for n, amount in enumerate(gross, 1):
        book["Payslip"].append([f"Person {n}", amount, amount * 0.2, amount * 0.8])
    book.save(target)
    notes.write_text(context)
    return source, target, notes


def statement_of(*payments, currency="USD"):
    """A statement with one dated deposit per (description, amount)."""
    header = f"{'Date':<15}{'Description':<55}{'Withdrawal':>14}{'Deposit':>14}{'Balance':>14}"
    lines, balance = [], 800
    for description, amount in payments:
        balance += amount
        lines.append(f"{'24 Feb 2026':<15}{description:<55}{'':>14}{f'{amount:.2f}':>14}{f'{balance:.2f}':>14}")
    return f"Currency : {currency}\n{header}\n" + "\n".join(lines) + "\n\n   Thank you for banking\n"


PAID_IN_FULL = statement_of(("REF TAYLOR REED", 100), ("REF MORGAN VALE", 50))
# Two descriptions of the same example. In it every reference has the full name and everyone paid exactly what
# was called, so the example cannot tell "amount called" from "amount paid", nor "full names only" from
# "partial names allowed". Only the description can.
PARTIAL_OK = ("The bank references are often cut short, so I match each payment to a partner by full name or, if "
              "needed, a partial name. I log what each partner actually paid according to the statement, and post two "
              "balancing lines per partner for that amount.")
FULL_ONLY = ("I match each payment to a partner by full name only; a reference with a partial name is not trusted and "
             "is left for follow-up. I log what each partner actually paid according to the statement, and post two "
             "balancing lines per partner for that amount.")
# The next month: Taylor's reference is cut short. (A short payment would stop the run instead: in the
# example every outstanding amount is 0, a value the learned checks hold to.)
NEXT_MONTH = statement_of(("REF TAYLOR REE", 100), ("REF MORGAN VALE", 50))


def paid_in_full(folder):
    """The fixture job with a completed example in which every partner paid exactly what was called."""
    source, target, bank, context = fixtures(folder)
    book = load_workbook(source)
    roster, stage, output = book["Roster"], book["Staging"], book["Import"]
    for row, received in ((3, 100), (4, 50)):
        roster.cell(row, 4, datetime(2026, 2, 24))
        roster.cell(row, 5, received)
        roster.cell(row, 6, 0)
    for reference, name, amount in (("REF TAYLOR REED", "Taylor Reed", 100), ("REF MORGAN VALE", "Morgan Vale", 50)):
        stage.append([datetime(2026, 2, 24), reference, amount, name])
    for entry, name, amount in [(1, "Taylor Reed", 100), (2, "Morgan Vale", 50)]:
        comment, paid = "Paid by " + name, datetime(2026, 2, 24)
        output.append(["Example fund", entry, 1, paid, "Cash received", "CASH-CODE", "Y", amount, comment, "Example account", "General", None, entry * 2 - 1])
        output.append(["Example fund", entry, 2, paid, "Due from partners", "RECEIVABLE-CODE", "N", amount, comment, None, "General", name, entry * 2])
    book.save(target)
    book.close()
    return source, target, bank, context
